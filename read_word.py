import docx
import sys

def read_word_document(file_path):
    try:
        doc = docx.Document(file_path)
        content = []
        
        # 提取段落内容
        for para in doc.paragraphs:
            if para.text.strip():
                content.append(para.text)
        
        # 提取表格内容
        for table in doc.tables:
            for row in table.rows:
                row_data = []
                for cell in row.cells:
                    if cell.text.strip():
                        row_data.append(cell.text.strip())
                if row_data:
                    content.append(" | ".join(row_data))
        
        return "\n".join(content)
    except Exception as e:
        return f"Error reading document: {str(e)}"

if __name__ == "__main__":
    file_path = r"C:\Users\ruancanling\Desktop\大数据基础课程考查题目及要求.docx"
    content = read_word_document(file_path)
    print(content)