"""工程导向型专属API：BOM多版本、ECO变更、WBS成本、设备、质量追溯、生产排程"""
import json
import io
import datetime
from typing import List, Optional
from fastapi import APIRouter, Depends, HTTPException, UploadFile, File
from sqlalchemy.orm import Session

from ..database import get_db
from .. import models
from ..app import make_response
from .notifications import push_notification

router = APIRouter()


# ==================== BOM多版本管理 ====================

@router.get("/bom-versions/")
def list_bom_versions(account_set_id: int = 1, product_id: Optional[int] = None, db: Session = Depends(get_db)):
    query = db.query(models.BOM).filter(models.BOM.account_set_id == account_set_id)
    if product_id:
        query = query.filter(models.BOM.product_id == product_id)
    boms = query.order_by(models.BOM.product_id, models.BOM.created_at.desc()).all()
    result = []
    for b in boms:
        result.append({
            "id": b.id,
            "product_id": b.product_id,
            "product_name": b.product.name if b.product else "",
            "version": b.version,
            "effective_date": str(b.effective_date) if b.effective_date else None,
            "end_date": str(b.end_date) if b.end_date else None,
            "status": b.status,
            "item_count": len(b.items),
            "items": [{"id": i.id, "material_name": i.material.name if i.material else "", "quantity": float(i.quantity), "unit": i.unit} for i in b.items],
            "created_at": str(b.created_at)
        })
    return result


@router.post("/bom-versions/")
def create_bom_version(data: dict, db: Session = Depends(get_db)):
    bom = models.BOM(
        account_set_id=data.get("account_set_id", 1),
        product_id=data["product_id"],
        version=data.get("version", "V1"),
        effective_date=datetime.date.fromisoformat(data["effective_date"]) if data.get("effective_date") else None,
        status=data.get("status", "ACTIVE")
    )
    db.add(bom)
    db.flush()
    for item in data.get("items", []):
        bi = models.BOMItem(
            bom_id=bom.id,
            material_id=item["material_id"],
            quantity=item["quantity"],
            unit=item.get("unit", "个"),
            scrap_rate=item.get("scrap_rate", 0),
            sequence=item.get("sequence", 0),
            level=item.get("level", 1)
        )
        db.add(bi)
    db.commit()
    db.refresh(bom)
    return {"id": bom.id, "version": bom.version}


@router.post("/bom-versions/{bom_id}/activate")
def activate_bom_version(bom_id: int, db: Session = Depends(get_db)):
    bom = db.query(models.BOM).filter(models.BOM.id == bom_id).first()
    if not bom:
        raise HTTPException(404, "BOM不存在")
    # 将同产品的其他版本置为INACTIVE
    db.query(models.BOM).filter(
        models.BOM.product_id == bom.product_id,
        models.BOM.id != bom_id
    ).update({"status": "INACTIVE"})
    bom.status = "ACTIVE"
    db.commit()
    return {"message": f"BOM版本 {bom.version} 已激活"}


# ==================== ECO工程变更单 ====================

@router.get("/eco-orders/")
def list_eco_orders(account_set_id: int = 1, status: Optional[str] = None, db: Session = Depends(get_db)):
    query = db.query(models.ECOChangeOrder).filter(models.ECOChangeOrder.account_set_id == account_set_id)
    if status:
        query = query.filter(models.ECOChangeOrder.status == status)
    orders = query.order_by(models.ECOChangeOrder.created_at.desc()).all()
    result = []
    for o in orders:
        result.append({
            "id": o.id,
            "eco_no": o.eco_no,
            "product_name": o.product_name,
            "old_version": o.old_version,
            "new_version": o.new_version,
            "change_type": o.change_type,
            "change_reason": o.change_reason,
            "change_content": json.loads(o.change_content) if o.change_content else [],
            "status": o.status,
            "requested_by": o.requested_by,
            "approved_by": o.approved_by,
            "effective_date": str(o.effective_date) if o.effective_date else None,
            "created_at": str(o.created_at)
        })
    return result


@router.post("/eco-orders/")
def create_eco_order(data: dict, db: Session = Depends(get_db)):
    eco_no = f"ECO-{datetime.datetime.now().strftime('%Y%m%d')}-{db.query(models.ECOChangeOrder).count() + 1:03d}"
    # 如果没有指定old_bom_id但有product_id，自动查找当前生效的BOM
    old_bom_id = data.get("old_bom_id")
    old_version = data.get("old_version", "")
    product_id = data.get("product_id")
    if not old_bom_id and product_id:
        active_bom = db.query(models.BOM).filter(
            models.BOM.product_id == product_id,
            models.BOM.status == "ACTIVE"
        ).first()
        if active_bom:
            old_bom_id = active_bom.id
            old_version = active_bom.version
    order = models.ECOChangeOrder(
        account_set_id=data.get("account_set_id", 1),
        eco_no=eco_no,
        product_id=product_id,
        product_name=data.get("product_name", ""),
        old_bom_id=old_bom_id,
        old_version=old_version,
        change_type=data.get("change_type", "material_substitute"),
        change_reason=data.get("change_reason", ""),
        change_content=json.dumps(data.get("change_content", []), ensure_ascii=False),
        status="DRAFT",
        requested_by=data.get("requested_by", "admin")
    )
    db.add(order)
    db.commit()
    db.refresh(order)
    return {"id": order.id, "eco_no": order.eco_no}


@router.post("/eco-orders/{order_id}/submit")
def submit_eco_order(order_id: int, db: Session = Depends(get_db)):
    order = db.query(models.ECOChangeOrder).filter(models.ECOChangeOrder.id == order_id).first()
    if not order:
        raise HTTPException(404, "ECO单不存在")
    order.status = "PENDING"
    db.commit()
    return {"message": "已提交审批"}


@router.post("/eco-orders/{order_id}/approve")
def approve_eco_order(order_id: int, data: dict = None, db: Session = Depends(get_db)):
    order = db.query(models.ECOChangeOrder).filter(models.ECOChangeOrder.id == order_id).first()
    if not order:
        raise HTTPException(404, "ECO单不存在")
    data = data or {}
    order.status = "APPROVED"
    order.approved_by = data.get("approved_by", "admin")
    order.approved_at = datetime.datetime.utcnow()
    order.effective_date = datetime.date.fromisoformat(data["effective_date"]) if data.get("effective_date") else datetime.date.today()
    db.commit()
    return {"message": "已审批通过"}


@router.post("/eco-orders/{order_id}/implement")
def implement_eco_order(order_id: int, db: Session = Depends(get_db)):
    """执行ECO：基于旧BOM创建新版本，标记旧版本失效"""
    order = db.query(models.ECOChangeOrder).filter(models.ECOChangeOrder.id == order_id).first()
    if not order:
        raise HTTPException(404, "ECO单不存在")
    if order.status != "APPROVED":
        raise HTTPException(400, "只有已审批的ECO单才能执行")
    # 如果没有指定旧BOM，自动查找该产品当前生效的BOM
    old_bom = None
    if order.old_bom_id:
        old_bom = db.query(models.BOM).filter(models.BOM.id == order.old_bom_id).first()
    if not old_bom and order.product_id:
        old_bom = db.query(models.BOM).filter(
            models.BOM.product_id == order.product_id,
            models.BOM.status == "ACTIVE"
        ).first()
    if not old_bom:
        raise HTTPException(400, "找不到对应的旧BOM，请先创建BOM并指定product_id")
    # 创建新版本BOM
    old_ver_num = int(old_bom.version.lstrip("Vv")) if old_bom.version.lstrip("Vv").isdigit() else 1
    new_ver = f"V{old_ver_num + 1}"
    new_bom = models.BOM(
        account_set_id=old_bom.account_set_id,
        product_id=old_bom.product_id,
        version=new_ver,
        effective_date=order.effective_date,
        status="ACTIVE"
    )
    db.add(new_bom)
    db.flush()
    # 复制旧BOM明细
    for item in old_bom.items:
        db.add(models.BOMItem(
            bom_id=new_bom.id,
            material_id=item.material_id,
            quantity=item.quantity,
            unit=item.unit,
            scrap_rate=item.scrap_rate,
            sequence=item.sequence,
            level=item.level
        ))
    # 应用变更内容
    changes = json.loads(order.change_content) if order.change_content else []
    for ch in changes:
        if ch.get("type") == "update_qty":
            for ni in new_bom.items:
                if ni.material_id == ch.get("material_id"):
                    ni.quantity = ch.get("new_qty", ni.quantity)
        elif ch.get("type") == "replace_material":
            for ni in new_bom.items:
                if ni.material_id == ch.get("old_material_id"):
                    ni.material_id = ch.get("new_material_id", ni.material_id)
        elif ch.get("type") == "add_item":
            db.add(models.BOMItem(
                bom_id=new_bom.id,
                material_id=ch.get("material_id"),
                quantity=ch.get("quantity", 1),
                unit=ch.get("unit", "个"),
                level=1
            ))
        elif ch.get("type") == "remove_item":
            for ni in list(new_bom.items):
                if ni.material_id == ch.get("material_id"):
                    db.delete(ni)
    # 旧版本失效
    old_bom.status = "INACTIVE"
    order.new_bom_id = new_bom.id
    order.new_version = new_ver
    order.status = "IMPLEMENTED"
    # 影响分析
    impact = {"affected_work_orders": 0, "affected_inventory": 0, "affected_purchase": 0}
    order.impact_analysis = json.dumps(impact, ensure_ascii=False)
    db.commit()
    return {"message": f"ECO已执行，BOM升级至 {new_ver}", "new_bom_id": new_bom.id, "new_version": new_ver}


# ==================== WBS项目管理 ====================

def recalc_project_costs(db: Session, account_set_id: int = 1):
    """自动归集项目成本（来源：已过账凭证 voucher_entries），并回写项目成本/净利润/毛利率。
    - 直接材料成本：借 1401 原材料（质检合格实际入库的采购材料）且挂项目的分录发生额
    - 制造费用分摊：借 5101 制造费用（车间共同费用，如水费电费）按各项目直接材料占比分摊
    - 项目成本 = 直接材料 + 分摊制造费用；净利润 = 合同额 - 项目成本；毛利率 = 净利润/合同额
    返回归集结果明细。
    """
    from sqlalchemy import func as _f
    # 1) 各项目直接材料成本（1401原材料借方=质检合格实际入库的采购材料；不含1402在途物资）
    mat_rows = db.query(
        models.VoucherEntry.project_id,
        _f.sum(models.VoucherEntry.debit)
    ).join(
        models.VoucherDB, models.VoucherEntry.voucher_id == models.VoucherDB.id
    ).filter(
        models.VoucherDB.status == models.VoucherStatus.POSTED,
        models.VoucherEntry.project_id.isnot(None),
        models.VoucherEntry.account_code == "1401",
    ).group_by(models.VoucherEntry.project_id).all()
    material_map = {pid: float(amt or 0) for pid, amt in mat_rows}

    # 2) 制造费用池（5101 借方，车间共同费用）
    overhead_pool = db.query(_f.sum(models.VoucherEntry.debit)).join(
        models.VoucherDB, models.VoucherEntry.voucher_id == models.VoucherDB.id
    ).filter(
        models.VoucherDB.status == models.VoucherStatus.POSTED,
        models.VoucherEntry.account_code == "5101",
    ).scalar() or 0
    overhead_pool = float(overhead_pool or 0)

    # 3) 分摊（按直接材料占比；无材料则不分摊）
    total_material = sum(material_map.values())
    projects = db.query(models.WBSProject).filter(
        models.WBSProject.account_set_id == account_set_id).all()
    detail = []
    for p in projects:
        mat = round(material_map.get(p.id, 0.0), 2)
        if total_material > 0:
            oh = round(overhead_pool * mat / total_material, 2)
        else:
            oh = 0.0
        total_cost = round(mat + oh, 2)
        contract = float(p.contract_amount or 0)
        net = round(contract - total_cost, 2)
        margin = round(net / contract * 100, 2) if contract > 0 else 0.0
        p.material_cost = mat
        p.overhead_cost = oh
        p.incurred_cost = total_cost
        p.net_profit = net
        p.gross_margin = margin
        detail.append({
            "project_id": p.id, "project_no": p.project_no,
            "project_name": p.project_name, "contract_amount": contract,
            "material_cost": mat, "overhead_cost": oh,
            "total_cost": total_cost, "net_profit": net, "gross_margin": margin,
        })
    db.commit()
    return {"overhead_pool": round(overhead_pool, 2),
            "total_material": round(total_material, 2),
            "projects": detail}


@router.get("/wbs-projects/recalc-costs")
def recalc_costs_endpoint(db: Session = Depends(get_db)):
    """手动触发：按财务凭证重新归集所有项目成本/净利润/毛利率"""
    try:
        return {"success": True, "data": recalc_project_costs(db)}
    except Exception as e:
        import traceback
        traceback.print_exc()
        return {"success": False, "message": str(e)}


@router.get("/wbs-projects/cost-report")
def project_cost_report(db: Session = Depends(get_db)):
    """项目成本明细表（报表层）：按成本要素列示 + 预算vs实际差异 + 超支预警。
    成本要素：直接材料(1401原材料，实际入库)、直接人工(5002)、外协费用(科目名含外协/委外)、制造费用(5101分摊)。
    """
    from sqlalchemy import func as _f
    # 先按最新凭证归集
    try:
        recalc_project_costs(db)
    except Exception:
        import traceback; traceback.print_exc()

    # 直接人工（5002 借方，挂项目）
    labor_rows = db.query(
        models.VoucherEntry.project_id, _f.sum(models.VoucherEntry.debit)
    ).join(models.VoucherDB, models.VoucherEntry.voucher_id == models.VoucherDB.id).filter(
        models.VoucherDB.status == models.VoucherStatus.POSTED,
        models.VoucherEntry.project_id.isnot(None),
        models.VoucherEntry.account_code == "5002",
    ).group_by(models.VoucherEntry.project_id).all()
    labor_map = {pid: float(amt or 0) for pid, amt in labor_rows}

    # 外协费用（科目名含"外协/委外"的借方，挂项目）
    out_rows = db.query(
        models.VoucherEntry.project_id, _f.sum(models.VoucherEntry.debit)
    ).join(models.VoucherDB, models.VoucherEntry.voucher_id == models.VoucherDB.id).filter(
        models.VoucherDB.status == models.VoucherStatus.POSTED,
        models.VoucherEntry.project_id.isnot(None),
        (models.VoucherEntry.account_name.like("%外协%") | models.VoucherEntry.account_name.like("%委外%")),
    ).group_by(models.VoucherEntry.project_id).all()
    out_map = {pid: float(amt or 0) for pid, amt in out_rows}

    projects = db.query(models.WBSProject).order_by(models.WBSProject.created_at.desc()).all()
    rows = []
    tot = {"contract": 0.0, "budget": 0.0, "material": 0.0, "labor": 0.0,
           "outsource": 0.0, "overhead": 0.0, "actual": 0.0, "profit": 0.0}
    for p in projects:
        material = float(p.material_cost or 0)
        labor = round(labor_map.get(p.id, 0.0), 2)
        outsource = round(out_map.get(p.id, 0.0), 2)
        overhead = float(p.overhead_cost or 0)
        actual = round(material + labor + outsource + overhead, 2)
        contract = float(p.contract_amount or 0)
        budget = float(p.budget_cost or 0)
        profit = round(contract - actual, 2)
        margin = round(profit / contract * 100, 2) if contract > 0 else 0.0
        usage_pct = round(actual / budget * 100, 1) if budget > 0 else 0.0  # 预算执行率
        budget_remain = round(budget - actual, 2)
        # 预警：超支红 / ≥80%预算橙 / 否则绿
        if budget > 0 and actual > budget:
            warn, warn_cn = "OVER", "已超支"
        elif budget > 0 and usage_pct >= 80:
            warn, warn_cn = "NEAR", "接近预算"
        elif budget > 0:
            warn, warn_cn = "OK", "预算内"
        else:
            warn, warn_cn = "NOBUDGET", "未设预算"
        rows.append({
            "project_id": p.id, "project_no": p.project_no, "project_name": p.project_name,
            "status": p.status,
            "contract_amount": contract, "budget_cost": budget,
            "material_cost": material, "labor_cost": labor,
            "outsource_cost": outsource, "overhead_cost": overhead,
            "actual_cost": actual, "net_profit": profit, "gross_margin": margin,
            "usage_pct": usage_pct, "budget_remain": budget_remain,
            "warn": warn, "warn_cn": warn_cn,
        })
        tot["contract"] += contract; tot["budget"] += budget
        tot["material"] += material; tot["labor"] += labor
        tot["outsource"] += outsource; tot["overhead"] += overhead
        tot["actual"] += actual; tot["profit"] += profit
    tot = {k: round(v, 2) for k, v in tot.items()}
    tot["gross_margin"] = round(tot["profit"] / tot["contract"] * 100, 2) if tot["contract"] > 0 else 0.0
    over_cnt = sum(1 for r in rows if r["warn"] == "OVER")
    near_cnt = sum(1 for r in rows if r["warn"] == "NEAR")
    return {
        "success": True,
        "data": {
            "rows": rows, "total": tot,
            "over_budget_count": over_cnt, "near_budget_count": near_cnt,
            "project_count": len(rows),
        },
    }


@router.get("/wbs-projects/")
def list_wbs_projects(account_set_id: int = 1, status: Optional[str] = None, db: Session = Depends(get_db)):
    # 每次查看项目前，先按最新财务凭证自动归集成本（幂等）
    try:
        recalc_project_costs(db, account_set_id)
    except Exception:
        import traceback
        traceback.print_exc()
    query = db.query(models.WBSProject).filter(models.WBSProject.account_set_id == account_set_id)
    if status:
        query = query.filter(models.WBSProject.status == status)
    projects = query.order_by(models.WBSProject.created_at.desc()).all()
    result = []
    for p in projects:
        # 计算完工百分比
        pct = float(p.progress_pct) if p.progress_pct else 0
        # 应确认收入 = 合同金额 × 完工百分比
        should_revenue = float(p.contract_amount) * pct / 100 if p.contract_amount else 0
        result.append({
            "id": p.id,
            "project_no": p.project_no,
            "project_name": p.project_name,
            "customer_name": p.customer_name or (p.customer.name if p.customer else ""),
            "currency": p.currency or "CNY",
            "delivery_mode": p.delivery_mode or "MTO",
            "contract_amount": float(p.contract_amount),
            "budget_cost": float(p.budget_cost),
            "incurred_cost": float(p.incurred_cost),
            "material_cost": float(p.material_cost or 0),
            "overhead_cost": float(p.overhead_cost or 0),
            "net_profit": float(p.net_profit or 0),
            "gross_margin": float(p.gross_margin or 0),
            "estimated_total_cost": float(p.estimated_total_cost),
            "progress_pct": pct,
            "recognized_revenue": float(p.recognized_revenue),
            "should_revenue": round(should_revenue, 2),
            "status": p.status,
            "manager": p.manager,
            "planned_start": str(p.planned_start) if p.planned_start else None,
            "planned_end": str(p.planned_end) if p.planned_end else None,
            "node_count": len(p.wbs_nodes),
            "deliverable_count": len(p.deliverables)
        })
    return result


# ==================== 项目合同 Excel 导出 / 模板 / 导入 ====================
_PROJECT_EXPORT_HEADERS = ["项目编号", "项目名称", "客户名称", "合同金额", "交付方式", "状态", "计划开始", "计划结束", "项目经理", "进度%", "备注"]
_DELIVERY_MODES = {"MTS", "MTO", "ATO", "ETO"}
_PROJECT_STATUSES = {"PLANNING", "EXECUTING", "PAUSED", "COMPLETED", "CLOSED"}


@router.get("/wbs-projects/export")
def export_wbs_projects(db: Session = Depends(get_db)):
    """导出全部项目合同（含状态/进度/日期等完整字段）"""
    from fastapi.responses import StreamingResponse
    import openpyxl
    projects = db.query(models.WBSProject).order_by(models.WBSProject.created_at.desc()).all()
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "项目合同"
    ws.append(_PROJECT_EXPORT_HEADERS)
    for p in projects:
        ws.append([
            p.project_no or "", p.project_name or "", p.customer_name or "",
            float(p.contract_amount or 0), p.delivery_mode or "MTO", p.status or "PLANNING",
            p.planned_start.strftime("%Y-%m-%d") if p.planned_start else "",
            p.planned_end.strftime("%Y-%m-%d") if p.planned_end else "",
            p.manager or "", float(p.progress_pct or 0), p.remark or "",
        ])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=projects.xlsx"})


@router.get("/wbs-projects/import-template")
def wbs_projects_import_template():
    """项目合同导入模板"""
    from fastapi.responses import StreamingResponse
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "项目合同"
    ws.append(_PROJECT_EXPORT_HEADERS)
    ws.append(["PRJ-20260901-001", "机器人搬运系统", "广州A公司", 1500000, "MTO", "PLANNING", "2026-09-10", "2026-12-31", "张工", 0, "含安装调试"])
    ws.append(["", "自动化产线二期", "深圳B公司", 800000, "ETO", "PLANNING", "", "", "李工", 0, ""])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=project_import_template.xlsx"})


@router.post("/wbs-projects/import")
async def import_wbs_projects(file: UploadFile = File(...), db: Session = Depends(get_db)):
    """批量导入项目合同：项目编号留空则自动生成，状态/交付方式自动校验转换"""
    import openpyxl
    content = await file.read()
    try:
        wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
    except Exception:
        raise HTTPException(400, "Excel文件无法读取")
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        raise HTTPException(400, "表格为空")
    headers = [str(h).strip() if h is not None else "" for h in rows[0]]

    def col(*names):
        for n in names:
            for idx, h in enumerate(headers):
                if h and n in h.replace(" ", ""):
                    return idx
        return -1
    c = {
        "no": col("项目编号"), "name": col("项目名称"), "cust": col("客户名称", "客户"),
        "amt": col("合同金额", "金额"), "mode": col("交付方式"), "status": col("状态"),
        "ps": col("计划开始"), "pe": col("计划结束"), "mgr": col("项目经理", "负责人"),
        "prog": col("进度"), "remark": col("备注"),
    }
    if c["name"] < 0:
        raise HTTPException(400, "缺少「项目名称」列")
    mode_map = {"mts": "MTS", "mto": "MTO", "ato": "ATO", "eto": "ETO",
                "备货": "MTS", "订单": "MTO", "装配": "ATO", "定制": "ETO"}
    status_map = {"planning": "PLANNING", "立项": "PLANNING", "立项中": "PLANNING",
                  "executing": "EXECUTING", "执行": "EXECUTING", "执行中": "EXECUTING",
                  "completed": "COMPLETED", "完成": "COMPLETED", "已完成": "COMPLETED",
                  "paused": "PAUSED", "暂停": "PAUSED", "closed": "CLOSED", "关闭": "CLOSED"}

    def _date(v):
        if v is None or str(v).strip() == "":
            return None
        if isinstance(v, datetime.datetime):
            return v.date()
        if isinstance(v, datetime.date):
            return v
        s = str(v).strip()
        for fmt in ("%Y-%m-%d", "%Y/%m/%d", "%Y.%m.%d", "%Y年%m月%d日"):
            try:
                return datetime.datetime.strptime(s, fmt).date()
            except ValueError:
                continue
        return None

    created, fail = 0, []
    base_cnt = db.query(models.WBSProject).count()
    for rn, row in enumerate(rows[1:], start=2):
        name = str(row[c["name"]]).strip() if c["name"] >= 0 and row[c["name"]] is not None else ""
        if not name or name in ("合计", "总计"):
            continue
        try:
            pno = str(row[c["no"]]).strip() if c["no"] >= 0 and row[c["no"]] is not None else ""
            if not pno:
                pno = f"PRJ-{datetime.datetime.now().strftime('%Y%m%d')}-{base_cnt + created + 1:03d}"
            if db.query(models.WBSProject).filter(models.WBSProject.project_no == pno).first():
                fail.append(f"第{rn}行: 项目编号 {pno} 已存在，跳过")
                continue
            mode = str(row[c["mode"]]).strip().upper() if c["mode"] >= 0 and row[c["mode"]] is not None else "MTO"
            if mode not in _DELIVERY_MODES:
                mode = mode_map.get(str(row[c["mode"]]).strip().lower() if row[c["mode"]] is not None else "", "MTO")
                mode = mode_map.get(str(row[c["mode"]]).strip() if row[c["mode"]] is not None else "", mode if mode in _DELIVERY_MODES else "MTO")
            status = str(row[c["status"]]).strip() if c["status"] >= 0 and row[c["status"]] is not None else "PLANNING"
            status = status_map.get(status.lower(), status_map.get(status, status if status in _PROJECT_STATUSES else "PLANNING"))
            amt = 0.0
            if c["amt"] >= 0 and row[c["amt"]] is not None:
                try:
                    amt = float(str(row[c["amt"]]).replace(",", "").replace("¥", "").replace("￥", ""))
                except ValueError:
                    fail.append(f"第{rn}行: 合同金额「{row[c['amt']]}」不是数字，按0处理")
            proj = models.WBSProject(
                account_set_id=1,
                project_no=pno,
                project_name=name,
                customer_name=str(row[c["cust"]]).strip() if c["cust"] >= 0 and row[c["cust"]] is not None else "",
                contract_amount=amt,
                delivery_mode=mode,
                status=status,
                planned_start=_date(row[c["ps"]]) if c["ps"] >= 0 else None,
                planned_end=_date(row[c["pe"]]) if c["pe"] >= 0 else None,
                manager=str(row[c["mgr"]]).strip() if c["mgr"] >= 0 and row[c["mgr"]] is not None else "",
                progress_pct=float(row[c["prog"]]) if c["prog"] >= 0 and row[c["prog"]] is not None and str(row[c["prog"]]).replace(".", "").isdigit() else 0,
                remark=str(row[c["remark"]]).strip() if c["remark"] >= 0 and row[c["remark"]] is not None else "",
            )
            db.add(proj)
            db.flush()
            if proj.status == "EXECUTING":
                _auto_create_production_task(db, proj)
            created += 1
        except Exception as e:
            fail.append(f"第{rn}行: {str(e)[:60]}")
    db.commit()
    msg = f"导入完成：成功{created}个项目"
    if fail:
        msg += f"，失败{len(fail)}条"
    return {"success": True, "created": created, "fail": len(fail), "fail_rows": fail[:10], "message": msg}


@router.post("/wbs-projects/")
def create_wbs_project(data: dict, db: Session = Depends(get_db)):
    # 优先使用前端传入的项目编号，否则自动生成
    proj_no = data.get("project_no") or f"PRJ-{datetime.datetime.now().strftime('%Y%m%d')}-{db.query(models.WBSProject).count() + 1:03d}"
    # 检查项目编号是否重复
    existing = db.query(models.WBSProject).filter(models.WBSProject.project_no == proj_no).first()
    if existing:
        raise HTTPException(400, f"项目编号 {proj_no} 已存在，请更换")
    project = models.WBSProject(
        account_set_id=data.get("account_set_id", 1),
        project_no=proj_no,
        project_name=data["project_name"],
        customer_id=data.get("customer_id"),
        customer_name=data.get("customer_name"),
        contract_amount=data.get("contract_amount", 0),
        currency=data.get("currency", "CNY"),
        delivery_mode=data.get("delivery_mode", "MTO"),
        budget_cost=data.get("budget_cost", 0),
        estimated_total_cost=data.get("estimated_total_cost", 0),
        status=data.get("status", "PLANNING"),
        manager=data.get("manager", ""),
        business_mode=data.get("business_mode", "B"),
        planned_start=datetime.date.fromisoformat(data["planned_start"]) if data.get("planned_start") else None,
        planned_end=datetime.date.fromisoformat(data["planned_end"]) if data.get("planned_end") else None
    )
    db.add(project)
    db.commit()
    db.refresh(project)
    # 合同生产后（状态为EXECUTING）自动创建生产派工任务
    if project.status == "EXECUTING":
        _auto_create_production_task(db, project)
    return {
        "id": project.id,
        "project_no": project.project_no,
        "project_name": project.project_name,
        "customer_name": project.customer_name,
        "contract_amount": float(project.contract_amount or 0),
        "currency": project.currency,
        "delivery_mode": project.delivery_mode
    }


def _auto_create_production_task(db: Session, project):
    """项目转EXECUTING时自动创建生产派工任务（合同签订→下发生产）"""
    # 避免重复创建
    existing = db.query(models.ProjectProductionTask).filter(
        models.ProjectProductionTask.project_id == project.id
    ).first()
    if existing:
        return existing
    task_no = f"PPT-{project.project_no}-{datetime.datetime.now().strftime('%H%M%S')}"
    task = models.ProjectProductionTask(
        project_id=project.id,
        task_no=task_no,
        status="CONTRACT_SIGNED",
        contract_signed_at=datetime.datetime.utcnow()
    )
    db.add(task)
    db.commit()
    return task


@router.put("/wbs-projects/{project_id}")
def update_wbs_project(project_id: int, data: dict, db: Session = Depends(get_db)):
    """更新WBS项目（含成本归集：incurred_cost）"""
    project = db.query(models.WBSProject).filter(models.WBSProject.id == project_id).first()
    if not project:
        raise HTTPException(404, "项目不存在")
    old_status = project.status
    for field in ["project_name", "customer_id", "customer_name", "contract_amount", "currency", "delivery_mode", "budget_cost",
                  "incurred_cost", "estimated_total_cost", "progress_pct",
                  "status", "manager", "remark", "planned_start", "planned_end",
                  "actual_start", "actual_end"]:
        if field in data:
            val = data[field]
            if field in ("planned_start", "planned_end", "actual_start", "actual_end") and val:
                val = datetime.date.fromisoformat(val)
            setattr(project, field, val)
    db.commit()
    db.refresh(project)
    # 状态从非EXECUTING变为EXECUTING时，自动创建生产派工任务
    if old_status != "EXECUTING" and project.status == "EXECUTING":
        _auto_create_production_task(db, project)
    return {
        "id": project.id,
        "project_no": project.project_no,
        "incurred_cost": float(project.incurred_cost or 0),
        "recognized_revenue": float(project.recognized_revenue or 0),
        "progress_pct": float(project.progress_pct or 0),
    }


@router.delete("/wbs-projects/{project_id}")
def delete_wbs_project(project_id: int, db: Session = Depends(get_db)):
    """删除WBS项目：级联清理所有关联数据"""
    project = db.query(models.WBSProject).filter(models.WBSProject.id == project_id).first()
    if not project:
        raise HTTPException(404, "项目不存在")
    # 1. 删除成本归集
    db.query(models.CostCollection).filter(models.CostCollection.project_id == project_id).delete()
    # 2. 删除收入确认
    db.query(models.RevenueRecognition).filter(models.RevenueRecognition.project_id == project_id).delete()
    # 3. 删除里程碑
    db.query(models.Milestone).filter(models.Milestone.project_id == project_id).delete()
    # 4. 删除技术委派加工单
    db.query(models.TechOutsourcingOrder).filter(models.TechOutsourcingOrder.project_id == project_id).delete()
    # 5. 删除技术分解
    db.query(models.TechnicalDecomposition).filter(models.TechnicalDecomposition.project_id == project_id).delete()
    # 6. 删除生产派工任务及其物料需求
    tasks = db.query(models.ProjectProductionTask).filter(models.ProjectProductionTask.project_id == project_id).all()
    for t in tasks:
        db.query(models.ProjectMaterialRequirement).filter(models.ProjectMaterialRequirement.task_id == t.id).delete()
    db.query(models.ProjectProductionTask).filter(models.ProjectProductionTask.project_id == project_id).delete()
    # 7. 删除交付物料
    db.query(models.ProjectDeliverable).filter(models.ProjectDeliverable.project_id == project_id).delete()
    # 8. 删除WBS节点
    db.query(models.WBSNode).filter(models.WBSNode.project_id == project_id).delete()
    # 9. 置空可空关联
    db.query(models.PurchaseOrder).filter(models.PurchaseOrder.project_id == project_id).update({"project_id": None})
    db.query(models.Contract).filter(models.Contract.project_id == project_id).update({"project_id": None})
    db.query(models.WorkOrderProcess).filter(models.WorkOrderProcess.project_id == project_id).update({"project_id": None})
    # 10. 删除项目本身
    db.delete(project)
    db.commit()
    return {"success": True, "message": "项目已删除"}


@router.post("/wbs-projects/{project_id}/accumulate-cost")
def accumulate_project_cost(project_id: int, data: dict, db: Session = Depends(get_db)):
    """项目成本归集：累加 incurred_cost，并自动生成项目成本凭证"""
    project = db.query(models.WBSProject).filter(models.WBSProject.id == project_id).first()
    if not project:
        raise HTTPException(404, "项目不存在")
    amount = float(data.get("amount", 0))
    if amount <= 0:
        raise HTTPException(400, "归集金额必须大于0")
    project.incurred_cost = float(project.incurred_cost or 0) + amount
    db.commit()

    # 🔥 自动触发项目成本归集凭证
    voucher_result = None
    try:
        from .auto_voucher import trigger_auto_voucher
        voucher_result = trigger_auto_voucher(
            db=db,
            doc_type="PROJECT_COST",
            doc_id=project.id,
            doc_no=f"{project.project_no}-COST-{int(amount)}",
            account_set_id=project.account_set_id or 1,
            force=True,
        )
    except Exception as ve:
        voucher_result = {"success": False, "error": str(ve)}

    return {
        "id": project.id,
        "incurred_cost": float(project.incurred_cost or 0),
        "voucher": voucher_result,
    }


@router.get("/wbs-projects/{project_id}/nodes")
def get_wbs_nodes(project_id: int, db: Session = Depends(get_db)):
    nodes = db.query(models.WBSNode).filter(models.WBSNode.project_id == project_id).order_by(models.WBSNode.sort_order).all()
    return [{
        "id": n.id,
        "node_code": n.node_code,
        "node_name": n.node_name,
        "parent_id": n.parent_id,
        "level": n.level,
        "budget_cost": float(n.budget_cost),
        "incurred_cost": float(n.incurred_cost),
        "work_order_id": n.work_order_id,
        "status": n.status
    } for n in nodes]


@router.post("/wbs-projects/{project_id}/nodes")
def add_wbs_node(project_id: int, data: dict, db: Session = Depends(get_db)):
    node = models.WBSNode(
        project_id=project_id,
        node_code=data["node_code"],
        node_name=data["node_name"],
        parent_id=data.get("parent_id"),
        level=data.get("level", 1),
        budget_cost=data.get("budget_cost", 0),
        status=data.get("status", "PLANNING")
    )
    db.add(node)
    db.commit()
    db.refresh(node)
    return {"id": node.id}


@router.post("/wbs-projects/{project_id}/calculate-revenue")
def calculate_revenue(project_id: int, db: Session = Depends(get_db)):
    """完工百分比法计算应确认收入，并自动生成收入确认凭证"""
    project = db.query(models.WBSProject).filter(models.WBSProject.id == project_id).first()
    if not project:
        raise HTTPException(404, "项目不存在")
    # 完工百分比 = 已发生成本 / 预计总成本
    if project.estimated_total_cost and float(project.estimated_total_cost) > 0:
        pct = float(project.incurred_cost) / float(project.estimated_total_cost) * 100
        project.progress_pct = min(round(pct, 2), 100)
    # 应确认收入 = 合同金额 × 完工百分比
    should_revenue = float(project.contract_amount) * float(project.progress_pct) / 100
    project.recognized_revenue = round(should_revenue, 2)
    db.commit()

    # 🔥 自动触发收入确认凭证（事件驱动）
    voucher_result = None
    if project.recognized_revenue and float(project.recognized_revenue) > 0:
        try:
            from .auto_voucher import trigger_auto_voucher
            voucher_result = trigger_auto_voucher(
                db=db,
                doc_type="REVENUE_CONFIRM",
                doc_id=project.id,
                doc_no=project.project_no,
                account_set_id=project.account_set_id or 1,
                force=True,
            )
        except Exception as ve:
            voucher_result = {"success": False, "error": str(ve)}

    return {
        "progress_pct": float(project.progress_pct),
        "recognized_revenue": float(project.recognized_revenue),
        "incurred_cost": float(project.incurred_cost),
        "estimated_total_cost": float(project.estimated_total_cost),
        "voucher": voucher_result,
    }


# ==================== 设备管理 ====================

@router.get("/equipments/")
def list_equipments(account_set_id: int = 1, status: Optional[str] = None, db: Session = Depends(get_db)):
    query = db.query(models.Equipment).filter(models.Equipment.account_set_id == account_set_id)
    if status:
        query = query.filter(models.Equipment.status == status)
    equips = query.order_by(models.Equipment.created_at.desc()).all()
    return [{
        "id": e.id,
        "equipment_code": e.equipment_code,
        "equipment_name": e.equipment_name,
        "category": e.category,
        "model": e.model,
        "manufacturer": e.manufacturer,
        "workshop": e.workshop,
        "status": e.status,
        "oee": float(e.oee),
        "availability": float(e.availability),
        "performance": float(e.performance),
        "quality_rate": float(e.quality_rate),
        "last_maintenance": str(e.last_maintenance) if e.last_maintenance else None,
        "next_maintenance": str(e.next_maintenance) if e.next_maintenance else None
    } for e in equips]


@router.post("/equipments/")
def create_equipment(data: dict, db: Session = Depends(get_db)):
    equip = models.Equipment(
        account_set_id=data.get("account_set_id", 1),
        equipment_code=data["equipment_code"],
        equipment_name=data["equipment_name"],
        category=data.get("category"),
        model=data.get("model"),
        manufacturer=data.get("manufacturer"),
        workshop=data.get("workshop"),
        status=data.get("status", "RUNNING"),
        oee=data.get("oee", 0),
        availability=data.get("availability", 0),
        performance=data.get("performance", 0),
        quality_rate=data.get("quality_rate", 0)
    )
    db.add(equip)
    db.commit()
    db.refresh(equip)
    return {"id": equip.id}


# ==================== 质量追溯 ====================

@router.get("/quality-traces/")
def list_quality_traces(account_set_id: int = 1, batch_no: Optional[str] = None, db: Session = Depends(get_db)):
    query = db.query(models.QualityTrace).filter(models.QualityTrace.account_set_id == account_set_id)
    if batch_no:
        query = query.filter(models.QualityTrace.batch_no == batch_no)
    traces = query.order_by(models.QualityTrace.created_at.desc()).all()
    return [{
        "id": t.id,
        "batch_no": t.batch_no,
        "product_name": t.product.name if t.product else "",
        "work_order_id": t.work_order_id,
        "process_step": t.process_step,
        "material_name": t.material.name if t.material else "",
        "qc_result": t.qc_result,
        "qc_quantity": float(t.qc_quantity),
        "defect_description": t.defect_description,
        "inspector": t.inspector,
        "inspect_date": str(t.inspect_date) if t.inspect_date else None,
        "supplier": t.supplier
    } for t in traces]


@router.post("/quality-traces/")
def create_quality_trace(data: dict, db: Session = Depends(get_db)):
    trace = models.QualityTrace(
        account_set_id=data.get("account_set_id", 1),
        batch_no=data["batch_no"],
        product_id=data.get("product_id"),
        work_order_id=data.get("work_order_id"),
        process_step=data.get("process_step"),
        material_id=data.get("material_id"),
        qc_result=data.get("qc_result", "PENDING"),
        qc_quantity=data.get("qc_quantity", 0),
        defect_description=data.get("defect_description"),
        inspector=data.get("inspector"),
        inspect_date=datetime.date.fromisoformat(data["inspect_date"]) if data.get("inspect_date") else None,
        supplier=data.get("supplier")
    )
    db.add(trace)
    db.commit()
    db.refresh(trace)
    return {"id": trace.id}


@router.get("/quality-traces/{batch_no}/trace")
def trace_batch(batch_no: str, db: Session = Depends(get_db)):
    """双向追溯：通过批次号查找所有相关质量记录"""
    traces = db.query(models.QualityTrace).filter(models.QualityTrace.batch_no == batch_no).all()
    if not traces:
        return {"batch_no": batch_no, "found": False, "upstream": [], "downstream": []}
    product_ids = set(t.product_id for t in traces if t.product_id)
    material_ids = set(t.material_id for t in traces if t.material_id)
    work_order_ids = set(t.work_order_id for t in traces if t.work_order_id)
    # 上游：相同物料的来料检验
    upstream = db.query(models.QualityTrace).filter(
        models.QualityTrace.material_id.in_(material_ids)
    ).all() if material_ids else []
    # 下游：相同产品的成品检验
    downstream = db.query(models.QualityTrace).filter(
        models.QualityTrace.product_id.in_(product_ids)
    ).all() if product_ids else []
    return {
        "batch_no": batch_no,
        "found": True,
        "current": [{
            "process_step": t.process_step,
            "qc_result": t.qc_result,
            "inspector": t.inspector,
            "inspect_date": str(t.inspect_date) if t.inspect_date else None
        } for t in traces],
        "upstream": [{
            "batch_no": u.batch_no,
            "material_name": u.material.name if u.material else "",
            "qc_result": u.qc_result,
            "supplier": u.supplier
        } for u in upstream],
        "downstream": [{
            "batch_no": d.batch_no,
            "product_name": d.product.name if d.product else "",
            "qc_result": d.qc_result,
            "work_order_id": d.work_order_id
        } for d in downstream]
    }


# ==================== 生产排程 ====================

@router.get("/production-schedules/")
def list_schedules(account_set_id: int = 1, db: Session = Depends(get_db)):
    schedules = db.query(models.ProductionSchedule).filter(
        models.ProductionSchedule.account_set_id == account_set_id
    ).order_by(models.ProductionSchedule.planned_start).all()
    return [{
        "id": s.id,
        "work_order_no": s.work_order_no,
        "product_name": s.product_name,
        "workshop": s.workshop,
        "equipment_id": s.equipment_id,
        "equipment_name": s.equipment.equipment_name if s.equipment else "",
        "process_step": s.process_step,
        "planned_start": str(s.planned_start) if s.planned_start else None,
        "planned_end": str(s.planned_end) if s.planned_end else None,
        "planned_qty": float(s.planned_qty),
        "completed_qty": float(s.completed_qty),
        "status": s.status,
        "priority": s.priority
    } for s in schedules]


@router.post("/production-schedules/")
def create_schedule(data: dict, db: Session = Depends(get_db)):
    schedule = models.ProductionSchedule(
        account_set_id=data.get("account_set_id", 1),
        work_order_id=data.get("work_order_id"),
        work_order_no=data.get("work_order_no"),
        product_name=data.get("product_name"),
        workshop=data.get("workshop"),
        equipment_id=data.get("equipment_id"),
        process_step=data.get("process_step"),
        planned_start=datetime.datetime.fromisoformat(data["planned_start"]) if data.get("planned_start") else None,
        planned_end=datetime.datetime.fromisoformat(data["planned_end"]) if data.get("planned_end") else None,
        planned_qty=data.get("planned_qty", 0),
        status=data.get("status", "PLANNED"),
        priority=data.get("priority", "normal")
    )
    db.add(schedule)
    db.commit()
    db.refresh(schedule)
    return {"id": schedule.id}


# ==================== 项目交付物料清单 ====================
@router.get("/wbs-projects/{project_id}/deliverables")
def list_deliverables(project_id: int, db: Session = Depends(get_db)):
    project = db.query(models.WBSProject).filter(models.WBSProject.id == project_id).first()
    if not project:
        raise HTTPException(404, "项目不存在")
    items = db.query(models.ProjectDeliverable).filter(
        models.ProjectDeliverable.project_id == project_id
    ).order_by(models.ProjectDeliverable.id).all()
    return [{
        "id": d.id,
        "project_id": d.project_id,
        "material_code": d.material_code or "",
        "material_name": d.material_name,
        "specification": d.specification or "",
        "quantity": float(d.quantity or 1),
        "unit": d.unit or "个",
        "unit_price": float(d.unit_price) if d.unit_price else None,
        "remark": d.remark or "",
        "source": d.source or "manual"
    } for d in items]


@router.post("/wbs-projects/{project_id}/deliverables")
def add_deliverable(project_id: int, data: dict, db: Session = Depends(get_db)):
    project = db.query(models.WBSProject).filter(models.WBSProject.id == project_id).first()
    if not project:
        raise HTTPException(404, "项目不存在")
    item = models.ProjectDeliverable(
        project_id=project_id,
        material_code=data.get("material_code"),
        material_name=data["material_name"],
        specification=data.get("specification"),
        quantity=data.get("quantity", 1),
        unit=data.get("unit", "个"),
        unit_price=data.get("unit_price"),
        remark=data.get("remark"),
        source="manual"
    )
    db.add(item)
    db.commit()
    db.refresh(item)
    return {"id": item.id, "material_name": item.material_name, "msg": "添加成功"}


@router.post("/wbs-projects/{project_id}/deliverables/batch")
def batch_add_deliverables(project_id: int, data: dict, db: Session = Depends(get_db)):
    """批量添加交付物料（Excel导入确认后调用）"""
    project = db.query(models.WBSProject).filter(models.WBSProject.id == project_id).first()
    if not project:
        raise HTTPException(404, "项目不存在")
    rows = data.get("items", [])
    if not rows:
        raise HTTPException(400, "没有要导入的数据")
    count = 0
    for row in rows:
        item = models.ProjectDeliverable(
            project_id=project_id,
            material_code=row.get("material_code"),
            material_name=row.get("material_name", "").strip(),
            specification=row.get("specification"),
            quantity=row.get("quantity", 1),
            unit=row.get("unit", "个"),
            unit_price=row.get("unit_price"),
            remark=row.get("remark"),
            source="excel"
        )
        if item.material_name:
            db.add(item)
            count += 1
    db.commit()
    return {"success": True, "imported": count, "msg": f"成功导入 {count} 条交付物料"}


@router.delete("/deliverables/{deliverable_id}")
def delete_deliverable(deliverable_id: int, db: Session = Depends(get_db)):
    item = db.query(models.ProjectDeliverable).filter(models.ProjectDeliverable.id == deliverable_id).first()
    if not item:
        raise HTTPException(404, "交付物料不存在")
    db.delete(item)
    db.commit()
    return {"success": True, "msg": "删除成功"}


@router.post("/wbs-projects/{project_id}/deliverables/preview-excel")
async def preview_deliverables_excel(project_id: int, file: UploadFile = File(...), db: Session = Depends(get_db)):
    """上传Excel文件，解析并返回预览数据（不入库）"""
    project = db.query(models.WBSProject).filter(models.WBSProject.id == project_id).first()
    if not project:
        raise HTTPException(404, "项目不存在")
    try:
        import openpyxl
    except ImportError:
        raise HTTPException(500, "服务器未安装openpyxl库")
    try:
        content = await file.read()
        wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
        ws = wb.active
        rows = []
        headers = []
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i == 0:
                headers = [str(c).strip() if c else f"列{j+1}" for j, c in enumerate(row)]
            else:
                row_data = {}
                for j, val in enumerate(row):
                    key = headers[j] if j < len(headers) else f"列{j+1}"
                    row_data[key] = str(val).strip() if val is not None else ""
                if any(row_data.values()):
                    rows.append(row_data)
        return {"success": True, "headers": headers, "rows": rows, "total": len(rows)}
    except Exception as e:
        raise HTTPException(400, f"Excel解析失败: {str(e)}")


@router.get("/wbs-projects/{project_id}/deliverables/template")
def download_deliverables_template(project_id: int, db: Session = Depends(get_db)):
    """下载交付物料Excel模板"""
    try:
        from fastapi.responses import StreamingResponse
        import openpyxl
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "交付物料清单"
        # 表头（与实际业务Excel对齐，支持智能映射）
        headers = ["项目号", "层级", "物料编码", "物料名称", "规格型号", "来源", "数量", "单位", "单价(元)", "备注"]
        ws.append(headers)
        # 示例行
        ws.append(["P-2026-09-001", "L1", "M-01", "机器人搬运系统", "含机器人+控制器+底座", "外购+外协", 1, "套", 220800, "示例数据，请删除后填写实际数据"])
        ws.append(["P-2026-09-001", "L2", "M-01-01", "六轴机器人", "发那科R-30iB Plus，20kg", "外购", 1, "台", 185000, ""])
        ws.append(["P-2026-09-001", "L2", "M-01-02", "机器人控制器", "发那科R-30iB配套", "外购", 1, "台", 32000, ""])
        # 调整列宽
        for i, h in enumerate(headers):
            ws.column_dimensions[openpyxl.utils.get_column_letter(i+1)].width = max(15, len(h)*2)
        # 输出
        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        return StreamingResponse(
            buf,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": "attachment; filename=deliverables_template.xlsx"}
        )
    except ImportError:
        raise HTTPException(500, "服务器未安装openpyxl库")


# ==================== 生产派工跨部门协同流程 ====================

# 状态流转定义
PRODUCTION_FLOW = [
    {"key": "CONTRACT_SIGNED", "label": "合同签订", "desc": "合同已签订，自动下发生产"},
    {"key": "PRODUCTION_RECEIVED", "label": "生产接收", "desc": "生产部门已接收任务"},
    {"key": "STEPS_DEFINED", "label": "工序定义", "desc": "生产步骤已确定"},
    {"key": "MATERIALS_UPLOADED", "label": "物料上传", "desc": "零件分类表已上传"},
    {"key": "BOM_CREATED", "label": "生成BOM表", "desc": "BOM表已自动生成"},
    {"key": "PROCUREMENT_DISPATCHED", "label": "下发采购", "desc": "标准件已下发采购部门"},
    {"key": "PURCHASING", "label": "采购中", "desc": "采购部门正在执行"},
    {"key": "COMPLETED", "label": "采购完成", "desc": "标准件已齐套入库"},
    {"key": "OUTSOURCING_DISPATCHED", "label": "委派加工", "desc": "自制件已委派加工"},
    {"key": "IN_PRODUCTION", "label": "生产中", "desc": "委外加工与生产执行中"},
]

# 任务状态 → 流程阶段索引（BOM生成/委派加工/生产中按实际数据推导）
_BASE_STATUS_IDX = {
    "CONTRACT_SIGNED": 0,
    "PRODUCTION_RECEIVED": 1,
    "STEPS_DEFINED": 2,
    "MATERIALS_UPLOADED": 3,
    "PROCUREMENT_DISPATCHED": 5,
    "PURCHASING": 6,
    "COMPLETED": 7,
}


@router.get("/wbs-projects/{project_id}/production-task")
def get_project_production_task(project_id: int, db: Session = Depends(get_db)):
    """获取项目的生产派工任务（含流转时间线）"""
    task = db.query(models.ProjectProductionTask).filter(
        models.ProjectProductionTask.project_id == project_id
    ).order_by(models.ProjectProductionTask.created_at.desc()).first()
    if not task:
        return {"exists": False, "flow": PRODUCTION_FLOW}
    steps = []
    if task.production_steps:
        try:
            steps = json.loads(task.production_steps)
        except Exception:
            steps = []
    reqs = db.query(models.ProjectMaterialRequirement).filter(
        models.ProjectMaterialRequirement.task_id == task.id
    ).all()
    # 展示阶段推导：生成BOM表 / 委派加工 / 生产中 按实际数据推进
    idx = _BASE_STATUS_IDX.get(task.status, 0)
    bom_created = any(r.source == "bom_import" for r in reqs)
    if bom_created and idx < 4:
        idx = 4
    outs = db.query(models.TechOutsourcingOrder).filter(
        models.TechOutsourcingOrder.project_id == project_id,
        models.TechOutsourcingOrder.status != "CANCELLED",
    ).all()
    if outs:
        if idx < 8:
            idx = 8
        if any(o.status in ("PROCESSING", "COMPLETED") for o in outs):
            idx = 9
    display_status = PRODUCTION_FLOW[idx]["key"]
    return {
        "exists": True,
        "task": {
            "id": task.id,
            "task_no": task.task_no,
            "status": task.status,
            "production_steps": steps,
            "production_owner": task.production_owner,
            "procurement_owner": task.procurement_owner,
            "contract_signed_at": str(task.contract_signed_at) if task.contract_signed_at else None,
            "production_received_at": str(task.production_received_at) if task.production_received_at else None,
            "steps_defined_at": str(task.steps_defined_at) if task.steps_defined_at else None,
            "materials_uploaded_at": str(task.materials_uploaded_at) if task.materials_uploaded_at else None,
            "procurement_dispatched_at": str(task.procurement_dispatched_at) if task.procurement_dispatched_at else None,
            "purchasing_completed_at": str(task.purchasing_completed_at) if task.purchasing_completed_at else None,
        },
        "materials": [{
            "id": r.id, "material_code": r.material_code, "material_name": r.material_name,
            "specification": r.specification, "quantity": float(r.quantity or 0),
            "unit": r.unit, "unit_price": float(r.unit_price or 0) if r.unit_price else None,
            "purchase_status": r.purchase_status, "source": r.source
        } for r in reqs],
        "flow": PRODUCTION_FLOW,
        "display_status": display_status,
    }


@router.post("/production-tasks/{task_id}/receive")
def receive_production_task(task_id: int, data: dict = None, db: Session = Depends(get_db)):
    """生产部门接收任务"""
    task = db.query(models.ProjectProductionTask).filter(models.ProjectProductionTask.id == task_id).first()
    if not task:
        raise HTTPException(404, "派工任务不存在")
    if task.status not in ("CONTRACT_SIGNED", "PRODUCTION_RECEIVED"):
        raise HTTPException(400, f"当前状态 {task.status} 不允许接收")
    task.status = "PRODUCTION_RECEIVED"
    task.production_received_at = datetime.datetime.utcnow()
    if data and data.get("production_owner"):
        task.production_owner = data["production_owner"]
    db.commit()
    return {"success": True, "msg": "生产部门已接收任务", "status": task.status}


@router.post("/production-tasks/{task_id}/steps")
def define_production_steps(task_id: int, data: dict, db: Session = Depends(get_db)):
    """生产部门定义生产步骤"""
    task = db.query(models.ProjectProductionTask).filter(models.ProjectProductionTask.id == task_id).first()
    if not task:
        raise HTTPException(404, "派工任务不存在")
    if task.status not in ("PRODUCTION_RECEIVED", "STEPS_DEFINED"):
        raise HTTPException(400, f"请先接收任务，当前状态 {task.status}")
    steps = data.get("steps", [])
    if not steps:
        raise HTTPException(400, "生产步骤不能为空")
    task.production_steps = json.dumps(steps, ensure_ascii=False)
    task.status = "STEPS_DEFINED"
    task.steps_defined_at = datetime.datetime.utcnow()
    if data.get("production_owner"):
        task.production_owner = data["production_owner"]
    db.commit()
    return {"success": True, "msg": "生产步骤已保存", "status": task.status}


@router.post("/production-tasks/{task_id}/materials/preview-excel")
async def preview_material_excel(task_id: int, file: UploadFile = File(...), db: Session = Depends(get_db)):
    """生产部门上传原材料Excel预览"""
    task = db.query(models.ProjectProductionTask).filter(models.ProjectProductionTask.id == task_id).first()
    if not task:
        raise HTTPException(404, "派工任务不存在")
    try:
        import openpyxl
    except ImportError:
        raise HTTPException(500, "服务器未安装openpyxl库")
    try:
        content = await file.read()
        wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
        ws = wb.active
        rows = []
        headers = []
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i == 0:
                headers = [str(c).strip() if c else f"列{j+1}" for j, c in enumerate(row)]
            else:
                row_data = {}
                for j, val in enumerate(row):
                    key = headers[j] if j < len(headers) else f"列{j+1}"
                    row_data[key] = str(val).strip() if val is not None else ""
                if any(row_data.values()):
                    rows.append(row_data)
        return {"success": True, "headers": headers, "rows": rows, "total": len(rows)}
    except Exception as e:
        raise HTTPException(400, f"Excel解析失败: {str(e)}")


@router.post("/production-tasks/{task_id}/materials/batch")
def batch_add_materials(task_id: int, data: dict, db: Session = Depends(get_db)):
    """批量导入原材料需求（生产部门上传Excel后）"""
    task = db.query(models.ProjectProductionTask).filter(models.ProjectProductionTask.id == task_id).first()
    if not task:
        raise HTTPException(404, "派工任务不存在")
    if task.status not in ("STEPS_DEFINED", "MATERIALS_UPLOADED"):
        raise HTTPException(400, f"请先定义生产步骤，当前状态 {task.status}")
    rows = data.get("items", [])
    if not rows:
        raise HTTPException(400, "没有要导入的数据")
    count = 0
    for row in rows:
        item = models.ProjectMaterialRequirement(
            task_id=task_id,
            material_code=row.get("material_code"),
            material_name=row.get("material_name", "").strip(),
            specification=row.get("specification"),
            quantity=row.get("quantity", 1),
            unit=row.get("unit", "个"),
            unit_price=row.get("unit_price"),
            remark=row.get("remark"),
            source="excel"
        )
        if item.material_name:
            db.add(item)
            count += 1
    task.status = "MATERIALS_UPLOADED"
    task.materials_uploaded_at = datetime.datetime.utcnow()
    db.commit()
    return {"success": True, "imported": count, "msg": f"成功导入 {count} 项原材料需求", "status": task.status}


@router.post("/production-tasks/{task_id}/materials")
def add_material_manual(task_id: int, data: dict, db: Session = Depends(get_db)):
    """手动添加单条原材料需求"""
    task = db.query(models.ProjectProductionTask).filter(models.ProjectProductionTask.id == task_id).first()
    if not task:
        raise HTTPException(404, "派工任务不存在")
    item = models.ProjectMaterialRequirement(
        task_id=task_id,
        material_code=data.get("material_code"),
        material_name=data.get("material_name", "").strip(),
        specification=data.get("specification"),
        quantity=data.get("quantity", 1),
        unit=data.get("unit", "个"),
        unit_price=data.get("unit_price"),
        remark=data.get("remark"),
        source="manual"
    )
    if not item.material_name:
        raise HTTPException(400, "物料名称不能为空")
    db.add(item)
    if task.status in ("STEPS_DEFINED", "MATERIALS_UPLOADED"):
        task.status = "MATERIALS_UPLOADED"
        task.materials_uploaded_at = datetime.datetime.utcnow()
    db.commit()
    return {"success": True, "id": item.id, "msg": "原材料需求已添加"}


@router.delete("/material-requirements/{req_id}")
def delete_material_requirement(req_id: int, db: Session = Depends(get_db)):
    item = db.query(models.ProjectMaterialRequirement).filter(models.ProjectMaterialRequirement.id == req_id).first()
    if not item:
        raise HTTPException(404, "原材料需求不存在")
    db.delete(item)
    db.commit()
    return {"success": True, "msg": "删除成功"}


@router.post("/production-tasks/{task_id}/dispatch-procurement")
def dispatch_to_procurement(task_id: int, data: dict = None, db: Session = Depends(get_db)):
    """下发到采购部门"""
    task = db.query(models.ProjectProductionTask).filter(models.ProjectProductionTask.id == task_id).first()
    if not task:
        raise HTTPException(404, "派工任务不存在")
    if task.status not in ("MATERIALS_UPLOADED", "PROCUREMENT_DISPATCHED"):
        raise HTTPException(400, f"请先上传原材料清单，当前状态 {task.status}")
    req_count = db.query(models.ProjectMaterialRequirement).filter(
        models.ProjectMaterialRequirement.task_id == task.id
    ).count()
    if req_count == 0:
        raise HTTPException(400, "请先添加原材料需求清单")
    task.status = "PROCUREMENT_DISPATCHED"
    task.procurement_dispatched_at = datetime.datetime.utcnow()
    if data and data.get("procurement_owner"):
        task.procurement_owner = data["procurement_owner"]
    db.commit()
    return {"success": True, "msg": "已下发到采购部门", "status": task.status}


# ==================== 采购需求池（采购部门入口）====================

@router.get("/procurement/pool")
def get_procurement_pool(status: Optional[str] = None, db: Session = Depends(get_db)):
    """采购需求池：所有已下发采购的派工任务+原材料需求"""
    query = db.query(models.ProjectProductionTask).filter(
        models.ProjectProductionTask.status.in_(["PROCUREMENT_DISPATCHED", "PURCHASING", "COMPLETED"])
    )
    if status:
        query = query.filter(models.ProjectProductionTask.status == status)
    tasks = query.order_by(models.ProjectProductionTask.procurement_dispatched_at.desc()).all()
    result = []
    for task in tasks:
        project = task.project
        reqs = db.query(models.ProjectMaterialRequirement).filter(
            models.ProjectMaterialRequirement.task_id == task.id
        ).all()
        result.append({
            "task_id": task.id,
            "task_no": task.task_no,
            "project_id": project.id if project else None,
            "project_no": project.project_no if project else "",
            "project_name": project.project_name if project else "",
            "status": task.status,
            "procurement_owner": task.procurement_owner,
            "dispatched_at": str(task.procurement_dispatched_at) if task.procurement_dispatched_at else None,
            "material_count": len(reqs),
            "total_amount": sum(float(r.quantity or 0) * float(r.unit_price or 0) for r in reqs if r.unit_price),
            "materials": [{
                "id": r.id, "material_code": r.material_code, "material_name": r.material_name,
                "specification": r.specification, "quantity": float(r.quantity or 0),
                "unit": r.unit, "unit_price": float(r.unit_price or 0) if r.unit_price else None,
                "purchase_status": r.purchase_status
            } for r in reqs]
        })
    return result


@router.put("/material-requirements/{req_id}/purchase-status")
def update_purchase_status(req_id: int, data: dict, db: Session = Depends(get_db)):
    """采购部门更新原材料采购状态"""
    item = db.query(models.ProjectMaterialRequirement).filter(models.ProjectMaterialRequirement.id == req_id).first()
    if not item:
        raise HTTPException(404, "原材料需求不存在")
    new_status = data.get("purchase_status")
    if new_status not in ("PENDING", "PURCHASING", "ARRIVED", "CANCELLED"):
        raise HTTPException(400, "无效的采购状态")
    item.purchase_status = new_status
    task = db.query(models.ProjectProductionTask).filter(models.ProjectProductionTask.id == item.task_id).first()
    if task:
        all_reqs = db.query(models.ProjectMaterialRequirement).filter(
            models.ProjectMaterialRequirement.task_id == task.id
        ).all()
        if all(r.purchase_status in ("ARRIVED", "CANCELLED") for r in all_reqs):
            task.status = "COMPLETED"
            task.purchasing_completed_at = datetime.datetime.utcnow()
        elif new_status in ("PURCHASING", "ARRIVED") and task.status == "PROCUREMENT_DISPATCHED":
            task.status = "PURCHASING"
    db.commit()
    return {"success": True, "msg": "采购状态已更新"}


@router.post("/procurement/purchase-all")
def purchase_all_from_pool(db: Session = Depends(get_db)):
    """一键采购：把采购需求池中所有「待采购」物料按项目生成采购订单草稿（DRAFT待确认）"""
    import uuid as _uuid
    from decimal import Decimal
    tasks = db.query(models.ProjectProductionTask).filter(
        models.ProjectProductionTask.status.in_(["PROCUREMENT_DISPATCHED", "PURCHASING"])).all()
    if not tasks:
        return make_response(False, None, "采购需求池为空，没有可采购的需求")
    by_project = {}
    for t in tasks:
        reqs = db.query(models.ProjectMaterialRequirement).filter(
            models.ProjectMaterialRequirement.task_id == t.id,
            models.ProjectMaterialRequirement.purchase_status == "PENDING").all()
        if reqs:
            by_project.setdefault(t.project_id or 0, []).extend((r, t) for r in reqs)
    if not by_project:
        return make_response(False, None, "没有「待采购」状态的物料需求")

    def _find_or_create_material(r):
        m = None
        if r.material_code:
            m = db.query(models.Material).filter(models.Material.code == r.material_code).first()
        if not m:
            m = db.query(models.Material).filter(models.Material.name == r.material_name).first()
        if not m:
            code = r.material_code or f"POOL-{_uuid.uuid4().hex[:6].upper()}"
            while db.query(models.Material).filter(models.Material.code == code).first():
                code = f"POOL-{_uuid.uuid4().hex[:6].upper()}"
            m = models.Material(account_set_id=1, name=r.material_name[:100], code=code,
                                unit=r.unit or "个", unit_price=float(r.unit_price or 0),
                                property=models.MaterialProperty.PURCHASE)
            db.add(m)
            db.flush()
        return m

    total_pos, total_items, po_nos, touched = 0, 0, [], set()
    today = datetime.date.today()
    for pid, pairs in by_project.items():
        po = models.PurchaseOrder(
            account_set_id=1,
            po_no=f"PO-{today.strftime('%Y%m%d')}-POOL-{_uuid.uuid4().hex[:4].upper()}",
            supplier_id=None, status="DRAFT", approval_status="PENDING",
            tax_rate=Decimal("13"), discount_amount=0,
            order_date=today, total_amount=0,
            project_id=pid or None, remark="采购需求池一键采购")
        db.add(po)
        db.flush()
        total = Decimal("0")
        for ln, (r, t) in enumerate(pairs, start=1):
            mat = _find_or_create_material(r)
            qty = Decimal(str(round(float(r.quantity or 1))))
            price = Decimal(str(float(r.unit_price or 0)))
            amt = (qty * price).quantize(Decimal("0.01"))
            db.add(models.PurchaseOrderItem(
                purchase_order_id=po.id, material_id=mat.id,
                quantity=int(qty), unit_price=price, line_no=ln,
                unit=r.unit or "", amount=amt,
                source_no=f"POOL#T{t.id}", received_qty=0,
                remark=r.material_code or ""))
            r.purchase_status = "PURCHASING"
            touched.add(t.id)
            total += amt
            total_items += 1
        po.total_amount = total
        po_nos.append(po.po_no)
        total_pos += 1
    for t in tasks:
        if t.id in touched and t.status == "PROCUREMENT_DISPATCHED":
            t.status = "PURCHASING"
            t.purchasing_completed_at = None
    db.commit()
    # 跨部门通知：一键采购结果上大厅新闻播报
    push_notification(db, f"采购需求池一键采购：{total_items} 项物料已转 {total_pos} 张采购订单",
                      "订单号：" + "、".join(po_nos[:5]) + (" 等" if len(po_nos) > 5 else "") + "，请审核并跟进供应商",
                      category="采购", source=",".join(po_nos[:3]))
    return make_response(True, {"orders": total_pos, "items": total_items, "po_nos": po_nos})
